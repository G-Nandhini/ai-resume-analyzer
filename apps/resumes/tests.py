import shutil
import tempfile
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document

from .forms import ResumeForm
from .models import Resume
from .services import (
    ResumeTextExtractionError,
    extract_resume_text,
    extract_text_from_docx,
)


class ResumeModelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='resumeuser',
            password='S3cure-Passphrase-2026',
        )

    def test_string_representation_returns_title(self):
        resume = Resume.objects.create(
            user=self.user,
            title='Backend Resume',
            file='resumes/backend.pdf',
        )

        self.assertEqual(str(resume), 'Backend Resume')

    def test_default_ordering_shows_newest_upload_first(self):
        older = Resume.objects.create(
            user=self.user,
            title='Older Resume',
            file='resumes/older.pdf',
        )
        newer = Resume.objects.create(
            user=self.user,
            title='Newer Resume',
            file='resumes/newer.pdf',
        )
        Resume.objects.filter(pk=older.pk).update(
            uploaded_at=timezone.now() - timedelta(days=1),
        )
        Resume.objects.filter(pk=newer.pk).update(uploaded_at=timezone.now())

        self.assertEqual(list(Resume.objects.values_list('title', flat=True)), [
            'Newer Resume',
            'Older Resume',
        ])


class ResumeFormTests(TestCase):
    def test_accepts_pdf_upload(self):
        form = ResumeForm(
            data={'title': 'PDF Resume'},
            files={'file': self.uploaded_file('resume.pdf', b'%PDF-1.4')},
        )

        self.assertTrue(form.is_valid(), form.errors)

    def test_accepts_docx_upload_case_insensitively(self):
        form = ResumeForm(
            data={'title': 'DOCX Resume'},
            files={'file': self.uploaded_file('resume.DOCX', b'docx data')},
        )

        self.assertTrue(form.is_valid(), form.errors)

    def test_rejects_unsupported_file_extension(self):
        form = ResumeForm(
            data={'title': 'Text Resume'},
            files={'file': self.uploaded_file('resume.txt', b'plain text')},
        )

        self.assertFalse(form.is_valid())
        self.assertIn('Only PDF and DOCX files are allowed.', form.errors['file'])

    def test_rejects_files_larger_than_five_mb(self):
        form = ResumeForm(
            data={'title': 'Large Resume'},
            files={'file': self.uploaded_file(
                'resume.pdf',
                b'a' * (ResumeForm.max_file_size + 1),
            )},
        )

        self.assertFalse(form.is_valid())
        self.assertIn('Resume file size must be 5MB or less.', form.errors['file'])

    def test_file_widget_limits_selectable_extensions(self):
        form = ResumeForm()

        self.assertEqual(form.fields['file'].widget.attrs['accept'], '.pdf,.docx')

    def uploaded_file(self, name, content):
        return SimpleUploadedFile(name, content)


class ResumeTextExtractionServiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='serviceuser',
            password='S3cure-Passphrase-2026',
        )

    def test_extract_text_from_docx_returns_paragraph_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / 'resume.docx'
            document = Document()
            document.add_paragraph('Candidate Name')
            document.add_paragraph('Python Developer')
            document.save(file_path)

            text = extract_text_from_docx(file_path)

        self.assertIn('Candidate Name', text)
        self.assertIn('Python Developer', text)

    @patch('apps.resumes.services.extract_text_from_pdf')
    def test_extract_resume_text_uses_pdf_extractor(self, mock_extract_pdf):
        mock_extract_pdf.return_value = 'PDF resume text'
        resume = Resume.objects.create(
            user=self.user,
            title='PDF Resume',
            file='resumes/candidate.pdf',
        )

        text = extract_resume_text(resume)

        self.assertEqual(text, 'PDF resume text')
        mock_extract_pdf.assert_called_once()

    def test_extract_resume_text_rejects_unsupported_extension(self):
        resume = Resume.objects.create(
            user=self.user,
            title='Text Resume',
            file='resumes/candidate.txt',
        )

        with self.assertRaises(ResumeTextExtractionError):
            extract_resume_text(resume)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ResumeViewTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        media_root = cls._overridden_settings['MEDIA_ROOT']
        super().tearDownClass()
        shutil.rmtree(media_root, ignore_errors=True)

    def setUp(self):
        self.user = User.objects.create_user(
            username='owner',
            password='S3cure-Passphrase-2026',
        )
        self.other_user = User.objects.create_user(
            username='other',
            password='S3cure-Passphrase-2026',
        )
        self.resume = Resume.objects.create(
            user=self.user,
            title='Owner Resume',
            file='resumes/owner.pdf',
        )

    def test_resume_list_requires_login(self):
        response = self.client.get(reverse('resume_list'))

        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('resume_list')}",
        )

    def test_resume_list_only_shows_current_users_resumes(self):
        other_resume = Resume.objects.create(
            user=self.other_user,
            title='Other Resume',
            file='resumes/other.pdf',
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse('resume_list'))

        self.assertContains(response, self.resume.title)
        self.assertNotContains(response, other_resume.title)

    @patch('apps.resumes.views.extract_resume_text')
    def test_upload_saves_resume_for_logged_in_user(self, mock_extract_resume_text):
        mock_extract_resume_text.return_value = 'Extracted candidate text'
        self.client.force_login(self.user)

        response = self.client.post(reverse('resume_upload'), {
            'title': 'Uploaded Resume',
            'file': SimpleUploadedFile('resume.pdf', b'%PDF-1.4'),
        })

        resume = Resume.objects.get(title='Uploaded Resume')
        self.assertEqual(resume.user, self.user)
        self.assertEqual(resume.extracted_text, 'Extracted candidate text')
        self.assertRedirects(response, reverse('resume_detail', kwargs={'id': resume.id}))

    @patch('apps.resumes.views.extract_resume_text')
    def test_upload_accepts_csrf_protected_post(self, mock_extract_resume_text):
        mock_extract_resume_text.return_value = 'Extracted candidate text'
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.user)

        get_response = csrf_client.get(reverse('resume_upload'))
        csrf_token = csrf_client.cookies['csrftoken'].value

        self.assertEqual(get_response.status_code, 200)
        self.assertIn('no-cache', get_response.headers['Cache-Control'])

        response = csrf_client.post(reverse('resume_upload'), {
            'csrfmiddlewaretoken': csrf_token,
            'title': 'CSRF Protected Resume',
            'file': SimpleUploadedFile('resume.pdf', b'%PDF-1.4'),
        })

        resume = Resume.objects.get(title='CSRF Protected Resume')
        self.assertRedirects(response, reverse('resume_detail', kwargs={'id': resume.id}))

    @patch('apps.resumes.views.extract_resume_text')
    def test_upload_keeps_file_when_text_extraction_fails(self, mock_extract_resume_text):
        mock_extract_resume_text.side_effect = ResumeTextExtractionError(
            'Your resume was uploaded, but we could not extract text from it.'
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse('resume_upload'), {
            'title': 'Unreadable Resume',
            'file': SimpleUploadedFile('resume.pdf', b'%PDF-1.4'),
        }, follow=True)

        resume = Resume.objects.get(title='Unreadable Resume')
        self.assertTrue(resume.file.storage.exists(resume.file.name))
        self.assertIsNone(resume.extracted_text)
        self.assertContains(
            response,
            'Your resume was uploaded, but we could not extract text from it.',
        )

    def test_detail_does_not_allow_access_to_another_users_resume(self):
        self.client.force_login(self.other_user)

        response = self.client.get(reverse('resume_detail', kwargs={'id': self.resume.id}))

        self.assertEqual(response.status_code, 404)

    def test_delete_requires_post_and_removes_owned_resume(self):
        self.client.force_login(self.user)

        get_response = self.client.get(reverse('resume_delete', kwargs={'id': self.resume.id}))
        post_response = self.client.post(reverse('resume_delete', kwargs={'id': self.resume.id}))

        self.assertEqual(get_response.status_code, 200)
        self.assertRedirects(post_response, reverse('resume_list'))
        self.assertFalse(Resume.objects.filter(pk=self.resume.pk).exists())
