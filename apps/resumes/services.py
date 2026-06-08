from pathlib import Path

import fitz
from docx import Document


class ResumeTextExtractionError(Exception):
    """Raised when text cannot be extracted from an uploaded resume."""


def extract_text_from_pdf(file_path):
    text_parts = []

    with fitz.open(file_path) as document:
        for page in document:
            text_parts.append(page.get_text('text'))

    return '\n'.join(text_parts).strip()


def extract_text_from_docx(file_path):
    document = Document(file_path)
    text_parts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    text_parts.append(cell_text)

    return '\n'.join(text_parts).strip()


def extract_resume_text(resume):
    file_path = resume.file.path
    extension = Path(resume.file.name).suffix.lower()

    try:
        if extension == '.pdf':
            text = extract_text_from_pdf(file_path)
        elif extension == '.docx':
            text = extract_text_from_docx(file_path)
        else:
            raise ResumeTextExtractionError('Only PDF and DOCX resumes can be processed.')
    except ResumeTextExtractionError:
        raise
    except Exception as exc:
        raise ResumeTextExtractionError(
            'Your resume was uploaded, but we could not extract text from it. '
            'Please make sure the file is readable and try again if needed.'
        ) from exc

    if not text:
        raise ResumeTextExtractionError(
            'Your resume was uploaded, but no readable text was found in the file.'
        )

    return text
